#!/usr/bin/env python3
"""Generate the final Hangzhou social-security inspection report from the retained DOCX template."""

from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path
import math
import shutil
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/zhuyanzhu/Documents/ChatGPT/自动巡检平台")
RUNTIME = ROOT / "runtime" / "hangzhou_shebao_20260801_27"
TEMPLATE = Path("/Users/zhuyanzhu/Desktop/巡检文件/杭州社保26年8月1日-24日数据报告.docx")
OUTPUT = ROOT / "杭州社保智能接待运营巡检报告（2026年8月1日-8月27日）.docx"

RECEPTION_TOTAL = 17_697
CONSULTATION_TOTAL = 37_239
AVG_ROUNDS = 2.1
REVIEW_TOTAL = 1_000
BAD_TOTAL = 334
GOOD_TOTAL = REVIEW_TOTAL - BAD_TOTAL
BAD_SESSIONS = 260
ACCURACY = GOOD_TOTAL / REVIEW_TOTAL * 100

TOP20 = [
    ("人工客服", 8483),
    ("人社小灵光", 1071),
    ("问候语", 871),
    ("网报新增回退申请表下载", 643),
    ("业务办理", 564),
    ("个人权益记录查询打印", 554),
    ("个体劳动者参保登记、停保登记", 372),
    ("跨省转移接续", 220),
    ("隶属关系变更", 187),
    ("职工参保登记表格下载", 181),
    ("人工客服工作时间", 144),
    ("停保该怎么操作", 132),
    ("当月新增回退", 128),
    ("社保缴费基数范围", 126),
    ("养老保险重复缴费怎么办", 108),
    ("申请补缴城镇职工社会保险费", 98),
    ("网上服务系统进入路径", 96),
    ("社保参保", 86),
    ("个人参保信息变更", 73),
    ("灵活就业人员参保条件", 71),
]

RECEPTION_DAILY = [
    136, 99, 908, 912, 904, 964, 965, 136, 70,
    948, 899, 965, 897, 969, 147, 91, 1007, 1039,
    873, 853, 845, 129, 67, 937, 844, 645, 448,
]

# The existing period chart exposes the first 24 daily consultation values.
# The final three dates are represented as one exact period total because
# separate daily values were not available in the retained materials.
CONSULTATION_DAILY = [
    177, 144, 1524, 1875, 2075, 2699, 2224, 196, 99,
    2087, 1818, 2213, 2145, 2367, 189, 135, 2136, 2340,
    1329, 1450, 1780, 177, 103, 2303,
]
CONSULTATION_LAST3_TOTAL = CONSULTATION_TOTAL - sum(CONSULTATION_DAILY)

BAD_CASES = [
    {
        "title": "案例一：新增参保时间选错，机器人未形成有效回复",
        "image": "case_bad_1_SS202608041629521189124715.png",
        "body": "市民咨询刚刚新增人员后因时间选择错误，想先减员再重新提交。机器人多次进入未形成有效回复的状态，仅展示推荐问题，未明确说明单位职工减员办理路径、可操作入口和注意事项，问题未闭环。",
        "action": "补充“新增时间填报错误、参保职工减少、重新提交”等关联问法，建立单位职工减员与信息更正的分支回复；首轮先确认办理主体，再给出线上入口、办理前提和注意事项，避免连续空回复。",
    },
    {
        "title": "案例二：灵活就业人员退休年龄咨询，回答未落到具体结论",
        "image": "case_bad_2_SS202608042252101994324906.png",
        "body": "市民提供出生年月、参保经历和“最早何时退休”等关键信息，机器人先后重复追问弹性提前或弹性延迟，未结合用户明确的“最早”诉求给出适用规则、年龄下限和办理前提。",
        "action": "增加“出生年月、灵活就业、最早退休年龄”组合问法的优先识别；先回答法定退休年龄和弹性提前退休边界，再补充最低缴费年限、待遇领取地等影响因素，并提示以经办机构核定结果为准。",
    },
    {
        "title": "案例三：失业金与8月缴费同时咨询，多意图未承接",
        "image": "case_bad_3_SS202608041800051821394547.png",
        "body": "市民同时咨询8月失业金领取、单位减员后何时可以领取，以及8月是否还需缴纳养老和医疗保险。机器人只围绕停保主体反复追问，未拆分失业待遇、养老缴费和医疗缴费三个问题。",
        "action": "增加多意图拆解流程：先分别回应失业金申领条件及时间，再说明减员后养老、医疗缴费衔接方式，最后给出线上办理入口和咨询渠道；对医保事项明确提示由医保经办机构负责，避免只回答其中一项。",
    },
]

GOOD_CASES = [
    {
        "title": "案例一：跨省转移接续，办理路径清晰完整",
        "image": "case_good_1_SS202608191636571501410525.png",
        "body": "机器人准确识别跨省转移接续事项，给出国家社会保险公共服务平台、掌上12333和电子社保卡三类办理渠道，并说明线上申请和进度查询路径，同时补充省内跨区域就业无需办理转移，信息完整且具有可执行性。",
    },
    {
        "title": "案例二：个人姓名变更，按办理主体区分路径",
        "image": "case_good_2_SS202608191638481295270544.png",
        "body": "机器人围绕“社保改名字”区分单位职工和个体劳动者两类场景，分别提供浙江政务服务网、浙里办及线下窗口办理方式，并列出材料提示和社保易窗线上申报渠道，形成了从判断主体到选择办理渠道的完整指引。",
    },
    {
        "title": "案例三：工伤认定，边界说明与后续指引到位",
        "image": "case_good_3_SS202608152051441142811875.png",
        "body": "机器人先明确工伤认定不属于社保中心工作范围，再提供工伤认定申请事项链接，并衔接工伤医疗费用、劳动能力鉴定和伤残待遇申领路径，既完成职责边界说明，又给出后续可执行步骤。",
    },
]


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ]
    path = next((p for p in candidates if Path(p).exists()), None)
    try:
        return ImageFont.truetype(path, size=size) if path else ImageFont.load_default()
    except OSError:
        return ImageFont.load_default()


def set_east_asia(run, name="Songti SC"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def style_run(run, size=14, bold=False, color=None):
    set_east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def replace_paragraph(paragraph, text, size=14, bold=False, align=None, color=None):
    paragraph.clear()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, color=color)


def add_body(document, text, bold_prefix=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        style_run(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        style_run(r)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_heading(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    style_run(r, size=16, bold=True)
    ppr = p._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    ppr.append(keep)
    return p


def add_case(document, item, kind):
    color = "E8A33B" if kind == "bad" else "53B495"
    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(7)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run(item["title"])
    style_run(r, bold=True, color=(255, 255, 255))
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    title._p.get_or_add_pPr().append(shd)

    image_p = document.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.add_run().add_picture(str(RUNTIME / item["image"]), width=Inches(6.35))
    image_p.paragraph_format.space_after = Pt(4)

    label = "问题表现：" if kind == "bad" else "优秀表现："
    add_body(document, label + item["body"], bold_prefix=label)
    if kind == "bad":
        add_body(document, "优化动作：" + item["action"], bold_prefix="优化动作：")


def set_cell_text(cell, text, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    style_run(r, size=10.5 if not header else 11, bold=header)
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), "D8EAF6" if header else "FFFFFF")
    cell.vertical_alignment = 1


def update_table(table):
    while len(table.rows) > len(TOP20) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(TOP20) + 1:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    set_cell_text(table.cell(0, 0), "知识库名称", header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), "会话量", header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (name, count) in enumerate(TOP20, 1):
        set_cell_text(table.cell(i, 0), name)
        set_cell_text(table.cell(i, 1), f"{count:,}", align=WD_ALIGN_PARAGRAPH.CENTER)


def crop_cards():
    src = Image.open(RUNTIME / "base_cards_current.png")
    # Current page coordinates for the first and second metric rows.
    top = src.crop((300, 305, 850, 445))
    middle = src.crop((300, 445, 850, 595))
    top.save(RUNTIME / "report_image1.png")
    middle.save(RUNTIME / "report_image2.png")


def draw_line_chart(path, values, title, color, labels, note=None):
    w, h = 1440, 520
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    f_title, f_axis, f_value = font(24, True), font(16), font(14)
    d.text((70, 22), title, fill=(40, 40, 40), font=f_title)
    left, right, top, bottom = 90, 1330, 85, 435
    vmax = max(values) * 1.16 if values else 1
    for tick in range(0, 5):
        y = bottom - (bottom - top) * tick / 4
        val = vmax * tick / 4
        d.line((left, y, right, y), fill=(225, 230, 238), width=1)
        d.text((18, y - 10), f"{int(val):,}", fill=(100, 100, 100), font=f_axis)
    points = []
    for i, val in enumerate(values):
        x = left + (right - left) * i / max(1, len(values) - 1)
        y = bottom - (bottom - top) * val / vmax
        points.append((x, y))
    if len(points) > 1:
        d.line(points, fill=color, width=5, joint="curve")
    for i, (x, y) in enumerate(points):
        d.ellipse((x - 4, y - 4, x + 4, y + 4), fill=color)
        if len(values) <= 27 or i in (0, len(values) - 1):
            d.text((x - 18, y - 26), f"{values[i]:,}", fill=(40, 40, 40), font=f_value)
        if i % max(1, math.ceil(len(values) / 12)) == 0 or i == len(values) - 1:
            d.text((x - 26, bottom + 12), labels[i], fill=(90, 90, 90), font=f_axis)
    if note:
        d.text((90, 470), note, fill=(120, 120, 120), font=font(15))
    im.save(path)


def draw_week_chart(path):
    start = date(2026, 8, 1)
    groups = [[] for _ in range(7)]
    for i, value in enumerate(RECEPTION_DAILY):
        groups[(start + timedelta(days=i)).weekday()].append(value)
    vals = [sum(x) for x in groups]
    labels = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
    draw_line_chart(path, vals, "周维度接待量分布", (76, 126, 214), labels)


def draw_unavailable_chart(path, title, subtitle):
    w, h = 1440, 520
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 30), title, fill=(40, 40, 40), font=font(26, True))
    d.rounded_rectangle((270, 145, 1170, 370), radius=12, outline=(205, 215, 230), width=3, fill=(250, 252, 255))
    d.text((520, 215), subtitle, fill=(105, 115, 130), font=font(28, True))
    d.text((430, 285), "本期保留原版配色与版式，不填入未经核验的数字", fill=(150, 155, 165), font=font(18))
    im.save(path)


def draw_donut_panel(path):
    w, h = 1440, 560
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.text((70, 28), "会话场景分析", fill=(40, 40, 40), font=font(26, True))
    titles = ["用户意图识别", "用户问句类型", "会话类型", "单/多轮概况"]
    colors = [(74, 137, 230), (93, 103, 196), (100, 109, 143), (196, 145, 111)]
    for i, (title, color) in enumerate(zip(titles, colors)):
        cx = 190 + i * 350
        cy = 285
        d.text((cx - 90, 100), title, fill=(60, 60, 60), font=font(21, True))
        d.ellipse((cx - 92, cy - 92, cx + 92, cy + 92), outline=color, width=30)
        d.ellipse((cx - 58, cy - 58, cx + 58, cy + 58), fill="white")
        d.text((cx - 58, cy - 7), "暂无有效", fill=(110, 115, 125), font=font(17))
        d.text((cx - 48, cy + 22), "分类数值", fill=(110, 115, 125), font=font(17))
    im.save(path)


def replace_media(docx_path, images):
    temp = docx_path.with_suffix(".media.tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        replacements = {f"word/media/image{i}.png": p.read_bytes() for i, p in enumerate(images, 1)}
        for info in source.infolist():
            target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    temp.replace(docx_path)


def update_image_heights(docx_path, images):
    doc = Document(docx_path)
    for shape, image in zip(doc.inline_shapes, images):
        with Image.open(image) as im:
            shape.height = int(shape.width * im.height / im.width)
    doc.save(docx_path)


def main():
    crop_cards()
    draw_line_chart(
        RUNTIME / "report_image3.png",
        RECEPTION_DAILY,
        "日维度接待量分布",
        (76, 126, 214),
        [(date(2026, 8, 1) + timedelta(days=i)).strftime("%m-%d") for i in range(27)],
    )
    draw_line_chart(
        RUNTIME / "report_image4.png",
        CONSULTATION_DAILY + [CONSULTATION_LAST3_TOTAL],
        "咨询量分布",
        (77, 180, 181),
        [(date(2026, 8, 1) + timedelta(days=i)).strftime("%m-%d") for i in range(24)] + ["8/25-27"],
        note=f"8月25日-8月27日合计 {CONSULTATION_LAST3_TOTAL:,} 次",
    )
    draw_donut_panel(RUNTIME / "report_image5.png")
    draw_week_chart(RUNTIME / "report_image6.png")
    draw_unavailable_chart(RUNTIME / "report_image7.png", "接待时段分布", "本期暂无可核验的分时分类数值")

    images = [RUNTIME / f"report_image{i}.png" for i in range(1, 8)]
    doc = Document(TEMPLATE)
    p = doc.paragraphs

    replace_paragraph(p[0], "杭州社保智能接待8月1日-8月27日数据分析报告", 24, color=(226, 88, 111), align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[1], "（26年8月1日-8月27日）", 16, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(
        p[2],
        "以下是2026年8月1日-8月27日“杭州社保”智能AI接待的数据分析汇报，本期从咨询数据分析、会话明细分析、市民咨询高频问题、待优化案例与线上优化动作、本期优秀案例及后续优化重点六个方面进行梳理。",
    )
    replace_paragraph(p[6], "杭州社保8月1日-8月27日咨询情况", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[8], "杭州社保8月1日-8月27日会话数分布", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[10], "杭州社保8月1日-8月27日咨询量分布", 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[11], "总体咨询情况。2026年8月1日-8月27日，“杭州社保”智能机器人接待量为17,697次，咨询量为37,239次，平均对话总轮次和有效平均对话轮次均为2.1轮。", bold=True)
    replace_paragraph(p[12], "当前AI意图识别情况。本期AI准确率未形成有效数值，人工巡检正确数和错误数均为0，未打标量为37,239条；本期人工复核1,000条记录，其中有效回复666条，发现待优化问题334条，复核准确率为66.6%。", bold=True)
    replace_paragraph(p[13], "当前咨询量分析。8月1日-8月27日接待量在67-1,039次之间波动，8月18日达到接待量峰值1,039次；咨询量集中在工作日，8月25日-8月27日合计3,654次。", bold=True)
    replace_paragraph(p[15], "8月1日-8月27日市民咨询高频问题", 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[17], "会话场景分析。市民咨询主要集中在人工客服、人社小灵光、问候语、网报新增回退申请表下载、业务办理、个人权益记录查询打印、个体劳动者参保登记与停保登记、跨省转移接续等事项。", bold=True)
    replace_paragraph(p[18], "高频问题分析。前20项高频问题中，人工客服相关咨询8,483次，人社小灵光1,071次，问候语871次，网报新增回退申请表下载643次，具体情况见下表。", bold=True)
    replace_paragraph(p[19], "接待高峰时间分析。接待量峰值为1,039次，出现在8月18日；从周维度看，工作日接待量明显高于周末，建议重点关注工作日高峰期间的首轮承接和复杂问题分流。", bold=True)
    replace_paragraph(p[21], "3、接待时段分析。本期分时分类数值未形成有效结果，后续将结合高峰时段运行表现，持续关注工作日白天的首轮回复、上下文承接和人工衔接稳定性。", bold=True)
    replace_paragraph(p[24], f"本期人工复核{REVIEW_TOTAL:,}条记录，发现待优化问题{BAD_TOTAL:,}条，覆盖{BAD_SESSIONS:,}个会话；有效回复{GOOD_TOTAL:,}条，复核准确率为{ACCURACY:.1f}%。问题主要表现为未形成有效回复、复杂问题未承接、多意图拆解不足及政策类问答未落到具体结论。", bold=False)
    replace_paragraph(p[25], "1. 补齐停保、减员、补缴、重复缴费、失业待遇及退休政策等高频复杂问法，增加组合条件和口语化问法。")
    replace_paragraph(p[26], "2. 建立多意图拆解和上下文承接机制，先分别回应用户问题，再补充办理入口、材料要求和责任边界。")
    replace_paragraph(p[27], "3. 对未形成有效回复、连续追问和政策边界不清场景设置兜底话术，并纳入高峰时段回归复测。")
    update_table(doc.tables[0])

    doc.add_page_break()
    add_heading(doc, "四、本期待优化案例")
    for item in BAD_CASES:
        add_case(doc, item, "bad")

    doc.add_page_break()
    add_heading(doc, "五、本期优秀案例")
    add_body(doc, "本期优秀案例体现出较好的场景识别、边界说明、办理路径指引和问题闭环能力，可沉淀为“识别场景—说明结论—提供入口—提示下一步”的通用答复结构。")
    for item in GOOD_CASES:
        add_case(doc, item, "good")

    add_heading(doc, "六、后续优化重点")
    add_body(doc, "1. 聚焦未形成有效回复的问题，优先补齐停保减员、重复缴费、失业待遇、医保衔接和退休政策等高频复杂场景。")
    add_body(doc, "2. 完善多轮对话状态记忆，避免重复追问；对同一用户的多个问题进行拆分回答，确保每个诉求都有明确结论。")
    add_body(doc, "3. 建立“问题收集—知识补齐—训练发布—回归复核”的闭环机制，持续跟踪工作日高峰时段首轮回复和人工衔接表现。")
    add_body(doc, "4. 将跨省转移、个人信息变更、工伤认定等优秀答复结构沉淀为可复用话术模板，提升相近问题的服务一致性。")

    doc.save(OUTPUT)
    replace_media(OUTPUT, images)
    update_image_heights(OUTPUT, images)
    replace_media(OUTPUT, images)
    print(OUTPUT)


if __name__ == "__main__":
    main()
