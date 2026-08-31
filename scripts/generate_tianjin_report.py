from __future__ import annotations

import io
import math
import shutil
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("/Users/zhuyanzhu/Documents/ChatGPT/自动巡检平台")
TEMPLATE = Path("/Users/zhuyanzhu/Desktop/郸城县大数据局8月1日-8月24日运营周报.docx")
OUTPUT = Path("/Users/zhuyanzhu/Desktop/天津河北分局8月1日-8月26日运营周报.docx")
PDF_OUTPUT = ROOT / "天津河北分局8月1日-8月26日运营周报.pdf"

ISSUE_IMAGE = ROOT / "tianjin_case_issue_image_backend.png"
JOB_IMAGE = ROOT / "tianjin_case_issue_job_backend.png"
DISPUTE_IMAGE = ROOT / "tianjin_case_issue_dispute_backend.png"
GOOD_IMAGE = ROOT / "tianjin_case_good_idcard_backend_top.png"
GOOD_APP_IMAGE = ROOT / "tianjin_case_good_5g_backend.png"

DAILY = [
    ("8/1", 195), ("8/2", 54), ("8/3", 152), ("8/4", 72),
    ("8/5", 54), ("8/6", 71), ("8/7", 80), ("8/8", 32),
    ("8/9", 62), ("8/10", 116), ("8/11", 92), ("8/12", 66),
    ("8/13", 84), ("8/14", 73), ("8/15", 58), ("8/16", 74),
    ("8/17", 119), ("8/18", 102), ("8/19", 66), ("8/20", 72),
    ("8/21", 97), ("8/22", 146), ("8/23", 130), ("8/24", 96),
    ("8/25", 83), ("8/26", 60),
]

NAVY = "#1F3A5F"
BLUE = "#4F6F91"
SLATE = "#71859A"
MUTED = "#8795A5"
RED = "#B64A4A"
GRAY = "#5B6570"
LIGHT = "#F2F5F8"
GRID = "#D5DDE5"


def get_font(size: int, bold: bool = False):
    candidates = [
        ("/System/Library/Fonts/Supplemental/Songti.ttc", 1 if bold else 6),
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
    ]
    for candidate in candidates:
        try:
            if isinstance(candidate, tuple):
                path, index = candidate
            else:
                path, index = candidate, 0
            return ImageFont.truetype(path, size=size, index=index)
        except Exception:
            pass
    return ImageFont.load_default()


def text(draw, xy, value, size=24, fill="#222222", bold=False, anchor=None):
    draw.text(xy, value, font=get_font(size, bold), fill=fill, anchor=anchor)


def make_base_chart(path: Path):
    im = Image.new("RGB", (1100, 460), "white")
    draw = ImageDraw.Draw(im)
    text(draw, (550, 28), "天津河北分局 AI文本机器人基础数据", 30, NAVY, True, "ma")
    text(draw, (550, 68), "统计周期：2026年8月1日-8月26日", 18, GRAY, False, "ma")
    cards = [
        ("会话总量", "2,306", BLUE),
        ("有效交互会话", "2,051", SLATE),
        ("累计有效轮次", "2,324", NAVY),
        ("2轮及以上", "183", MUTED),
        ("0有效轮次", "255", RED),
    ]
    x0, y0, w, h, gap = 38, 130, 195, 126, 18
    for i, (label, value, color) in enumerate(cards):
        x = x0 + i * (w + gap)
        draw.rounded_rectangle((x, y0, x + w, y0 + h), radius=8, fill=LIGHT, outline=GRID, width=2)
        draw.rectangle((x, y0, x + 8, y0 + h), fill=color)
        text(draw, (x + 22, y0 + 25), label, 20, GRAY)
        text(draw, (x + 22, y0 + 76), value, 34, color, True)
    text(draw, (40, 305), "期间运行概览", 23, NAVY, True)
    text(draw, (40, 350), "日均会话量 88.7 条", 21)
    text(draw, (310, 350), "峰值：8月1日 195 条", 21)
    text(draw, (650, 350), "低点：8月8日 32 条", 21)
    text(draw, (40, 405), "说明：本图统计基于巡检页会话列表，识别方式与回复标签用于分布分析，不换算为人工准确率。", 17, GRAY)
    im.save(path)


def make_daily_chart(path: Path):
    im = Image.new("RGB", (1100, 350), "white")
    draw = ImageDraw.Draw(im)
    text(draw, (550, 24), "每日会话量趋势", 28, NAVY, True, "ma")
    left, top, right, bottom = 70, 70, 1060, 290
    max_value = 210
    for tick in (0, 50, 100, 150, 200):
        y = bottom - tick / max_value * (bottom - top)
        draw.line((left, y, right, y), fill=GRID, width=1)
        text(draw, (left - 12, y), str(tick), 14, GRAY, anchor="rm")
    bar_w = (right - left) / len(DAILY) * 0.68
    step = (right - left) / len(DAILY)
    for i, (day, value) in enumerate(DAILY):
        x = left + i * step + (step - bar_w) / 2
        y = bottom - value / max_value * (bottom - top)
        color = RED if value == max(v for _, v in DAILY) else BLUE
        draw.rounded_rectangle((x, y, x + bar_w, bottom), radius=3, fill=color)
        text(draw, (x + bar_w / 2, y - 8), str(value), 13, color, True, "ms")
        if i % 2 == 0:
            text(draw, (x + bar_w / 2, bottom + 12), day, 13, GRAY, anchor="ma")
    draw.line((left, bottom, right, bottom), fill=GRAY, width=1)
    text(draw, (right, 320), "红色为本期峰值", 16, GRAY, anchor="ra")
    im.save(path)


def make_detail_chart(path: Path):
    im = Image.new("RGB", (1100, 330), "white")
    draw = ImageDraw.Draw(im)
    text(draw, (550, 24), "会话明细结构分析", 28, NAVY, True, "ma")
    sections = [
        ("有效对话轮次", [("1轮", 1868, BLUE), ("2轮及以上", 183, SLATE), ("0轮", 255, RED)]),
        ("识别方式", [("关键词", 1534, NAVY), ("AI知识库", 620, SLATE), ("图片参与", 57, MUTED)]),
        ("回复标签", [("直接回复相关", 1397, BLUE), ("未标注", 525, RED), ("静默相关", 284, MUTED)]),
    ]
    max_values = [1868, 1534, 1397]
    for si, (title, items) in enumerate(sections):
        x = 35 + si * 355
        text(draw, (x, 80), title, 20, NAVY, True)
        base_y = 125
        for j, (label, value, color) in enumerate(items):
            y = base_y + j * 58
            text(draw, (x, y), label, 17, GRAY)
            draw.rounded_rectangle((x + 112, y + 2, x + 300, y + 22), radius=4, fill="#E8EEF5")
            width = 180 * value / max_values[si]
            draw.rounded_rectangle((x + 112, y + 2, x + 112 + width, y + 22), radius=4, fill=color)
            text(draw, (x + 310, y + 11), f"{value:,}", 16, color, True, "lm")
    text(draw, (35, 305), "注：组合识别/组合标签已归入对应的“相关”统计项，分布合计可能包含多标签记录。", 16, GRAY)
    im.save(path)


def clear_paragraph(p):
    p._p.clear_content()


def style_run(run, size=14, bold=False, color="222222", font_name="Songti SC"):
    color = color.lstrip("#")
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(
    p,
    value,
    size=14,
    bold=False,
    color="222222",
    align=None,
    font_name="Songti SC",
):
    clear_paragraph(p)
    run = p.add_run(value)
    style_run(run, size, bold, color, font_name)
    if align is not None:
        p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0.74) if align != WD_ALIGN_PARAGRAPH.CENTER else Cm(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35


def add_body(doc, value, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if bold_prefix and value.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, 14, True)
        r2 = p.add_run(value[len(bold_prefix):])
        style_run(r2, 14)
    else:
        r = p.add_run(value)
        style_run(r, 14)
    return p


def add_heading(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(value)
    style_run(r, 16, True, NAVY, "Songti SC")
    return p


def add_caption(doc, value):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(value)
    style_run(r, 10, False, GRAY)
    return p


def add_image(doc, path: Path, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def replace_media(docx_path: Path, generated: dict[str, Path]):
    temp = Path(tempfile.mkdtemp(prefix="tianjin_docx_"))
    rewritten = temp / "report.docx"
    with ZipFile(docx_path, "r") as zin, ZipFile(rewritten, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in generated:
                data = generated[item.filename].read_bytes()
            zout.writestr(item, data)
    shutil.copyfile(rewritten, docx_path)
    shutil.rmtree(temp, ignore_errors=True)


def main():
    chart_base = ROOT / "tianjin_report_base.png"
    chart_daily = ROOT / "tianjin_report_daily.png"
    chart_detail = ROOT / "tianjin_report_detail.png"
    make_base_chart(chart_base)
    make_daily_chart(chart_daily)
    make_detail_chart(chart_detail)

    doc = Document(str(TEMPLATE))
    p = doc.paragraphs
    set_para(p[1], "“天津河北分局”AI运营周报", size=28, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, font_name="Songti SC")
    set_para(
        p[3],
        "以下是2026年8月1日-8月26日“天津河北分局”AI文本机器人运行的详细周报，我们将从咨询数据、会话明细、线上优化动作及优秀案例四个维度进行分析，以反映机器人线上运行情况并形成可执行的优化动作。",
        size=14,
    )
    set_para(p[4], "一、咨询数据分析", size=16, bold=True, font_name="Songti SC")
    set_para(p[6], "咨询情况", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(p[8], "接待量分布", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(
        p[9],
        "总体咨询情况。2026年8月1日-8月26日，天津河北分局巡检页共记录2306条会话，日均会话量为88.7条；其中2051条会话存在有效对话轮次，累计有效对话轮次2324次。",
        size=14,
    )
    set_para(
        p[10],
        "2、当前AI准确率。本期巡检页面未提供可直接用于本周期的人工正确、错误标注总量，因此本报告不将识别方式或回复标签分布换算为AI准确率，避免形成失真的准确率结论。",
        size=14,
    )
    set_para(
        p[11],
        "3、当前咨询量分析。日趋势显示，8月1日会话量最高，为195条；8月3日、8月22日、8月23日分别为152条、146条、130条，构成本期高峰区间；8月8日最低，为32条。",
        size=14,
    )
    set_para(p[12], "二、会话明细分析", size=16, bold=True, font_name="Songti SC")
    set_para(p[14], "市民关心的问题", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_para(
        p[15],
        "1、会话结构分析。本期共有1868条会话为1轮，占81.01%；183条会话达到2轮及以上，占7.94%；255条会话有效对话轮次为0，占11.06%。多轮会话占比仍有提升空间，适合重点关注追问承接和上下文保持。",
        size=14,
    )
    set_para(
        p[16],
        "2、识别与回复标签分析。识别方式以关键词1534条、AI知识库620条为主；图片参与识别57条。回复标签中直接回复相关1397条、未标注525条、静默相关284条、FAQ_RAG相关122条。上述统计反映系统处理路径分布，不等同人工准确率。",
        size=14,
    )
    set_para(p[18], "三、线上优化AI动作", size=16, bold=True, font_name="Songti SC")
    set_para(
        p[19],
        "1、建立非警务事项分流机制。针对“燃气停供”类问题被直接引导拨打110的情况，补充燃气、物业、水电等非警务事项识别规则，优先提供对应服务热线、物业报修和属地主管部门渠道，并明确110适用边界。",
        size=14,
    )
    set_para(
        p[20],
        "2、强化多轮上下文保持与意图复核。针对招聘诈骗会话在用户补充现场细节后误跳转“银行取现报备”的情况，增加历史意图继承、候选意图冲突检测和低置信度二次确认，避免诈骗类型在多轮对话中漂移。",
        size=14,
    )
    set_para(
        p[21],
        "3、完善图片识别失败兜底。针对“未成功识别图片信息”场景，增加图片重传提示、关键信息文字补充入口和人工渠道引导；同时把图片识别失败纳入专项抽样，区分图片格式、清晰度和内容类型定位原因。",
        size=14,
    )
    set_para(p[23], "四、本期优秀案例", size=16, bold=True, font_name="Songti SC")
    set_para(
        p[24],
        "1、身份证到期换证咨询。会话ID：SS202608101707331933858388。市民询问身份证10月份到期是否可以携带旧证到派出所换领，机器人明确说明换领条件、所需材料、河北分局十个户籍派出所的地址和电话，并补充周一至周日9时至17时的对外服务时间；市民随后回复“好的谢谢”，问题得到完整回应。",
        size=14,
    )
    clear_paragraph(p[25])
    clear_paragraph(p[26])
    add_image(doc, GOOD_IMAGE, 15.8)
    add_caption(doc, "优秀案例截图：会话ID SS202608101707331933858388")
    add_body(doc, "案例价值：该会话围绕“换证条件—所需材料—办理地点—联系方式—服务时间”组织答案，信息完整且具备直接办理价值，可作为户政类办事咨询的标准回复样本。")
    add_body(
        doc,
        "2、可疑软件风险识别。会话ID：SS202608101032291184397508。市民反映“全民5G北斗卫星导航”软件要求填写银行卡和身份证信息，机器人识别为电信网络诈骗，明确提示立即卸载、不要填写敏感信息并通过官方应用商店下载软件；市民进一步确认未泄露信息并完成卸载，机器人继续给出国家反诈中心APP和来电预警建议，形成了风险识别、阻断操作和后续防护的连续回应。",
    )
    add_image(doc, GOOD_APP_IMAGE, 15.8)
    add_caption(doc, "优秀案例截图：会话ID SS202608101032291184397508")
    add_body(doc, "案例价值：该会话能够根据用户的实际处置进展调整回复，既说明风险原因，又给出卸载、核查和后续防护动作，具有较好的问题解决完整度。")
    add_heading(doc, "五、回复不恰当案例及优化动作")
    add_body(
        doc,
        "1、图片识别失败。会话ID：SS202608261819291984585694。市民发送图片后，机器人仅返回“未成功识别图片信息”，未进一步引导用户补充文字、重新上传或转入可处理渠道。优化动作：增加失败原因提示和下一步操作按钮式话术，至少提供“重新上传图片、描述图片内容、转人工咨询”三类出口。",
    )
    add_image(doc, ISSUE_IMAGE, 15.8)
    add_caption(doc, "问题案例截图：会话ID SS202608261819291984585694")
    add_body(
        doc,
        "2、招聘诈骗误跳银行取现报备。会话ID：SS202608260958371561959839。首轮回复能够识别招聘诈骗并提示停止转账，市民继续补充工作群、冒充HR、实地查看办公场所等具体特征后，机器人却转为银行行业取现报备链接，出现多轮上下文丢失和意图分支错误。优化动作：保存首轮诈骗类型作为主意图；后续回复先复核“招聘诈骗/冒充HR/求职群诱导”标签，再决定是否进入报警、止损或线索收集分支。",
    )
    add_image(doc, JOB_IMAGE, 15.8)
    add_caption(doc, "问题案例截图：会话ID SS202608260958371561959839")
    add_body(
        doc,
        "3、非银行纠纷仍重复跳转银行报备。会话ID：SS202608212018371431686522。市民先表达希望民警帮助解决问题，机器人直接返回银行行业取现报备链接；市民明确补充“不是银行的事，是和朋友的纠纷问题”后，机器人仍重复返回同一链接，未完成问题澄清、纠纷分流或人工求助引导。优化动作：增加“非银行事项”否定意图识别和重复回复拦截；对纠纷类诉求先询问事件地点、是否存在人身或财产损害，再分流至派出所、110或其他责任部门。",
    )
    add_image(doc, DISPUTE_IMAGE, 15.8)
    add_caption(doc, "问题案例截图：会话ID SS202608212018371431686522")
    add_heading(doc, "六、后续训练优化")
    add_body(doc, "1、围绕非警务事项、招聘诈骗、图片识别失败等问题建立专项测试集，进行单轮和多轮回归验证。")
    add_body(doc, "2、补充天津河北区户政、燃气及其他高频公共服务的办理条件、联系方式、时间和分流规则，提升答案的可执行性。")
    add_body(doc, "3、针对0有效轮次、静默、未标注和AI_NO_REFERENCE类记录建立抽样巡检机制，形成问题发现、知识库调整和复测闭环。")

    doc.core_properties.title = "天津河北分局AI运营周报（2026年8月1日-8月26日）"
    doc.core_properties.subject = "天津河北分局文本机器人运营数据分析"
    doc.core_properties.author = "百应"
    doc.core_properties.comments = "基于巡检页期间会话数据生成"
    doc.save(str(OUTPUT))

    replace_media(
        OUTPUT,
        {
            "word/media/image2.png": chart_base,
            "word/media/image3.png": chart_daily,
            "word/media/image4.png": chart_detail,
        },
    )
    print(OUTPUT)


if __name__ == "__main__":
    main()
