#!/usr/bin/env python3
"""Build the Hangzhou social-security August 2026 inspection report."""

from copy import deepcopy
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path("/Users/zhuyanzhu/Documents/ChatGPT/自动巡检平台")
RUNTIME = ROOT / "runtime" / "hangzhou_shebao_20260801_27"
TEMPLATE = Path("/Users/zhuyanzhu/Desktop/巡检文件/杭州社保26年8月1日-24日数据报告.docx")
OUTPUT = ROOT / "杭州社保智能接待运营巡检报告（2026年8月1日-8月27日）.docx"


TOP20 = [
    ("人工客服", 8483),
    ("人社小灵光", 1071),
    ("问候语", 871),
    ("网报新增回退申请表下载", 643),
    ("业务办理", 564),
    ("个人权益记录查询打印", 554),
    ("个体劳动者参保登记、停保登记", 372),
    ("跨省转移接续", 220),
    ("隶属关系变更", 187),
    ("职工参保登记表格下载", 181),
    ("人工客服工作时间", 144),
    ("停保该怎么操作", 132),
    ("当月新增回退", 128),
    ("社保缴费基数范围", 126),
    ("养老保险重复缴费怎么办", 108),
    ("申请补缴城镇职工社会保险费", 98),
    ("网上服务系统进入路径", 96),
    ("社保参保", 86),
    ("个人参保信息变更", 73),
    ("灵活就业人员参保条件", 71),
]


CASES_BAD = [
    {
        "title": "案例一：新增参保时间选错，机器人未形成有效回复",
        "image": "case_bad_1_SS202608041629521189124715.png",
        "body": "市民咨询刚刚新增人员后因时间选择错误，想先减员再重新提交。机器人多次进入未形成有效回复的状态，仅展示推荐问题，未明确说明单位职工减员办理路径、可操作入口和注意事项，问题未闭环。",
        "action": "补充“新增时间填报错误/参保职工减少”关联问法，建立单位职工减员与信息更正的分支回复；首轮直接给出浙江政务服务网办理入口及办理前提，无法确认主体时先用一个问题完成主体识别，避免连续空回复。",
    },
    {
        "title": "案例二：灵活就业人员退休年龄咨询，回答未落到具体结论",
        "image": "case_bad_2_SS202608042252101994324906.png",
        "body": "市民提供出生年月、参保经历和“最早何时退休”等关键信息，机器人先后重复追问弹性提前或弹性延迟，未结合用户明确的“最早”诉求给出适用规则、年龄下限和办理前提。",
        "action": "增加“出生年月+灵活就业+最早退休年龄”组合问法的优先识别；先回答法定退休年龄和弹性提前退休的边界，再补充最低缴费年限、待遇领取地等影响因素，并提示以经办机构核定结果为准。",
    },
    {
        "title": "案例三：失业金与8月缴费同时咨询，多意图未承接",
        "image": "case_bad_3_SS202608041800051821394547.png",
        "body": "市民同时咨询8月失业金领取、单位减员后何时可以领取，以及8月是否还需缴纳养老和医疗保险。机器人只围绕停保主体反复追问，未拆分失业待遇、养老缴费和医疗缴费三个问题。",
        "action": "增加多意图拆解流程：先分别确认失业金申领条件及时间，再说明减员后养老、医疗缴费衔接方式，最后给出线上办理入口和咨询渠道；对医保事项明确提示由医保经办机构负责，避免只回答其中一项。",
    },
]


CASES_GOOD = [
    {
        "title": "案例一：跨省转移接续，办理路径清晰完整",
        "image": "case_good_1_SS202608191636571501410525.png",
        "body": "机器人准确识别跨省转移接续事项，给出国家社会保险公共服务平台、掌上12333和电子社保卡三类办理渠道，并说明线上申请和进度查询路径，同时补充省内跨区域就业无需办理转移，信息完整且具有可执行性。",
    },
    {
        "title": "案例二：个人姓名变更，按办理主体区分路径",
        "image": "case_good_2_SS202608191638481295270544.png",
        "body": "机器人围绕“社保改名字”区分单位职工和个体劳动者两类场景，分别提供浙江政务服务网、浙里办及线下窗口办理方式，并列出材料提示和社保易窗线上申报渠道，形成了从判断主体到选择办理渠道的完整指引。",
    },
    {
        "title": "案例三：工伤认定，边界说明与后续指引到位",
        "image": "case_good_3_SS202608152051441142811875.png",
        "body": "机器人先明确工伤认定不属于社保中心工作范围，再提供工伤认定申请事项链接，并衔接工伤医疗费用、劳动能力鉴定和伤残待遇申领路径，既完成职责边界说明，又给出后续可执行步骤。",
    },
]


def set_east_asia(run, name="宋体"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_run_style(run, *, name="宋体", size=14, bold=False, color=None):
    set_east_asia(run, name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def replace_paragraph(paragraph, text, *, name="宋体", size=14, bold=False, color=None, align=None):
    paragraph.clear()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_style(run, name=name, size=size, bold=bold, color=color)
    return paragraph


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    ppr = paragraph._p.get_or_add_pPr()
    if keep_lines:
        el = ppr.find(qn("w:keepLines"))
        if el is None:
            ppr.append(OxmlElement("w:keepLines"))
    if keep_next:
        el = ppr.find(qn("w:keepNext"))
        if el is None:
            ppr.append(OxmlElement("w:keepNext"))


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_section_heading(document, text, fill=None, color=(0, 0, 0)):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_style(run, size=16, bold=True, color=color)
    if fill:
        shade_paragraph(p, fill)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_body(document, text, *, bold_prefix=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_style(r1, size=14, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_style(r2, size=14)
    else:
        r = p.add_run(text)
        set_run_style(r, size=14)
    return p


def add_case(document, item, kind):
    fill = "E8A33B" if kind == "bad" else "53B495"
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(item["title"])
    set_run_style(r, size=14, bold=True, color=(255, 255, 255))
    shade_paragraph(p, fill)
    set_paragraph_keep(p, keep_next=True)

    img = RUNTIME / item["image"]
    ip = document.add_paragraph()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ir = ip.add_run()
    ir.add_picture(str(img), width=Inches(6.35))
    ip.paragraph_format.space_after = Pt(4)
    set_paragraph_keep(ip, keep_next=True)

    label = "问题表现：" if kind == "bad" else "优秀表现："
    add_body(document, label + item["body"], bold_prefix=label)
    if kind == "bad":
        add_body(document, "优化动作：" + item["action"], bold_prefix="优化动作：")


def set_cell_text(cell, text, *, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    set_run_style(r, size=10.5 if not header else 11, bold=header)
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), "D8EAF6" if header else "FFFFFF")
    cell.vertical_alignment = 1


def update_high_frequency_table(table):
    while len(table.rows) > len(TOP20) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(TOP20) + 1:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    set_cell_text(table.cell(0, 0), "知识库名称", header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), "会话量", header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, (name, count) in enumerate(TOP20, start=1):
        set_cell_text(table.cell(idx, 0), name)
        set_cell_text(table.cell(idx, 1), count, align=WD_ALIGN_PARAGRAPH.CENTER)


def crop_visuals():
    crops = {
        "visual_cards_current.png": ("base_cards_current.png", (250, 235, 1395, 700)),
        "visual_reception_current.png": ("base_charts_current.png", (250, 190, 1395, 690)),
        "visual_period_current.png": ("dashboard_base_full.png", (250, 255, 1395, 685)),
        "visual_activity_current.png": ("base_lower_current.png", (250, 120, 1395, 620)),
        "visual_detail_current.png": ("dashboard_base_current.png", (250, 250, 1395, 700)),
        "visual_reception_lower_current.png": ("base_charts_current.png", (250, 185, 1395, 700)),
        "visual_summary_current.png": ("base_cards_current.png", (250, 250, 1395, 690)),
    }
    paths = []
    for out_name, (src_name, box) in crops.items():
        src = RUNTIME / src_name
        out = RUNTIME / out_name
        with Image.open(src) as im:
            im.crop(box).save(out, format="PNG")
        paths.append(out)
    return paths


def replace_media(docx_path, images):
    temp = docx_path.with_suffix(".media.tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as dst:
        mapping = {f"word/media/image{i}.png": p.read_bytes() for i, p in enumerate(images, start=1)}
        for info in src.infolist():
            dst.writestr(info, mapping.get(info.filename, src.read(info.filename)))
    temp.replace(docx_path)


def update_image_heights(docx_path, images):
    doc = Document(docx_path)
    for shape, image in zip(doc.inline_shapes, images):
        with Image.open(image) as im:
            shape.height = int(shape.width * im.height / im.width)
    doc.save(docx_path)


def main():
    visuals = crop_visuals()
    doc = Document(TEMPLATE)
    p = doc.paragraphs

    replace_paragraph(p[0], "杭州社保智能接待8月1日-8月27日数据分析报告", name="宋体", size=24, color=(226, 88, 111), align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[1], "（26年8月1日-8月27日）", name="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(
        p[2],
        "以下是2026年8月1日-8月27日“杭州社保”智能AI接待的数据分析汇报，本期从咨询数据分析、会话明细分析、市民咨询高频问题、待优化案例与线上优化动作、本期优秀案例及后续优化重点六个方面进行梳理。",
        name="宋体",
        size=14,
    )
    replace_paragraph(p[6], "杭州社保8月1日-8月27日咨询情况", name="宋体", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[8], "杭州社保8月1日-8月27日会话数分布", name="宋体", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[10], "杭州社保8月1日-8月27日咨询量分布", name="宋体", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[11], "总体咨询情况。2026年8月1日-8月27日，“杭州社保”智能机器人接待量为17,697次，咨询量为37,239次，平均对话总轮次和有效平均对话轮次均为2.1轮。", name="宋体", size=14, bold=True)
    replace_paragraph(p[12], "当前AI意图识别情况。本期AI准确率未形成有效数值，人工巡检正确数和错误数均为0，未打标量为37,239条。后续以逐条复核结果作为回答质量改进依据。", name="宋体", size=14, bold=True)
    replace_paragraph(p[13], "当前咨询量分析。8月1日-8月27日每日接待量在67-1,039次之间波动，8月18日达到峰值1,039次；工作日接待量明显高于周末，日间工作时段为主要服务时段。", name="宋体", size=14, bold=True)
    replace_paragraph(p[15], "8月1日-8月27日市民咨询高频问题", name="宋体", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[17], "会话场景分析。市民咨询主要集中在人工客服、人社小灵光、问候语、网报新增回退申请表下载、业务办理、个人权益记录查询打印、个体劳动者参保登记与停保登记、跨省转移接续等事项。", name="宋体", size=14, bold=True)
    replace_paragraph(p[18], "高频知识库分析。会话量较高的内容集中在人工客服、人社小灵光、问候语及多项社保办理事项，前20项具体情况见下表。", name="宋体", size=14, bold=True)
    replace_paragraph(p[19], "接待高峰时间分析。日接待量峰值为1,039次，出现在8月18日；工作日咨询活跃度高于周末，建议重点关注工作日高峰时段的首轮承接和多轮追问效果。", name="宋体", size=14, bold=True)
    replace_paragraph(p[21], "3、接待时段分析。结合本期运行表现，服务压力主要集中在工作日白天时段，应优先保障高峰期间知识召回、上下文承接和转人工衔接稳定。", name="宋体", size=14, bold=True)

    replace_paragraph(p[23], "待优化案例与线上优化动作", name="宋体", size=16)
    replace_paragraph(
        p[24],
        "本次复核导出聊天记录1,000条，涉及556个会话；发现待优化问题334条，覆盖260个会话。按复核记录计算，有效回复666条，准确率为66.6%。问题主要表现为未形成有效回复、复杂问题未承接、多意图拆解不足、退休政策回答未落到具体结论、图片或链接内容无法识别等。",
        name="宋体",
        size=14,
    )
    replace_paragraph(p[25], "1. 补齐停保、减员、补缴、重复缴费、失业待遇等高频复杂问法，增加组合条件和口语化问法。", name="宋体", size=14)
    replace_paragraph(p[26], "2. 建立多意图拆解和上下文承接机制，先分别回应用户问题，再补充办理入口、材料要求和责任边界。", name="宋体", size=14)
    replace_paragraph(p[27], "3. 对未形成有效回复、图片识别失败和连续追问场景设置兜底话术，并纳入高峰时段回归复测。", name="宋体", size=14)

    table = doc.tables[0]
    update_high_frequency_table(table)

    doc.add_page_break()
    add_section_heading(doc, "四、本期待优化案例")
    for item in CASES_BAD:
        add_case(doc, item, "bad")

    doc.add_page_break()
    add_section_heading(doc, "五、本期优秀案例")
    add_body(doc, "本期优秀案例体现出较好的场景识别、边界说明、办理路径指引和问题闭环能力，可沉淀为“识别场景—说明结论—提供入口—提示下一步”的通用答复结构。")
    for item in CASES_GOOD:
        add_case(doc, item, "good")

    add_section_heading(doc, "六、后续优化重点")
    add_body(doc, "1. 聚焦未形成有效回复的问题，优先补齐停保减员、重复缴费、失业待遇、医保衔接、退休政策等高频复杂场景。")
    add_body(doc, "2. 完善多轮对话状态记忆，避免重复追问；对同一用户的多个问题进行拆分回答，确保每个诉求都有明确结论。")
    add_body(doc, "3. 建立“问题收集—知识补齐—训练发布—回归复核”的闭环机制，持续跟踪高峰时段首轮回复和转人工衔接表现。")
    add_body(doc, "4. 将跨省转移、个人信息变更、工伤认定等优秀答复结构沉淀为可复用话术模板，提升相近问题的服务一致性。")

    doc.save(OUTPUT)
    replace_media(OUTPUT, visuals)
    update_image_heights(OUTPUT, visuals)
    replace_media(OUTPUT, visuals)
    print(OUTPUT)


if __name__ == "__main__":
    main()
