#!/usr/bin/env python3
"""Rebuild the Hangzhou report using the Dancheng screenshot presentation style."""

from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path
import math
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/zhuyanzhu/Documents/ChatGPT/自动巡检平台")
RUNTIME = ROOT / "runtime" / "hangzhou_shebao_20260801_27"
TEMPLATE = Path("/Users/zhuyanzhu/Desktop/郸城县大数据局8月1日-8月24日运营周报.docx")
OUTPUT = ROOT / "杭州社保智能接待运营巡检报告（2026年8月1日-8月27日）.docx"

RECEPTION_DAILY = [136, 99, 908, 912, 904, 964, 965, 136, 70, 948, 899, 965, 897, 969, 147, 91, 1007, 1039, 873, 853, 845, 129, 67, 937, 844, 645, 448]
CONSULTATION_DAILY = [177, 144, 1524, 1875, 2075, 2699, 2224, 196, 99, 2087, 1818, 2213, 2145, 2367, 189, 135, 2136, 2340, 1329, 1450, 1780, 177, 103, 2303]
CONSULTATION_LAST3_TOTAL = 37239 - sum(CONSULTATION_DAILY)

TOP20 = [
    ("人工客服", 8483), ("人社小灵光", 1071), ("问候语", 871),
    ("网报新增回退申请表下载", 643), ("业务办理", 564),
    ("个人权益记录查询打印", 554), ("个体劳动者参保登记、停保登记", 372),
    ("跨省转移接续", 220), ("隶属关系变更", 187),
    ("职工参保登记表格下载", 181), ("人工客服工作时间", 144),
    ("停保该怎么操作", 132), ("当月新增回退", 128),
    ("社保缴费基数范围", 126), ("养老保险重复缴费怎么办", 108),
    ("申请补缴城镇职工社会保险费", 98), ("网上服务系统进入路径", 96),
    ("社保参保", 86), ("个人参保信息变更", 73), ("灵活就业人员参保条件", 71),
]

BAD_CASES = [
    ("案例一：新增参保时间选错，机器人未形成有效回复", "case_bad_1_SS202608041629521189124715.png",
     "市民咨询刚刚新增人员后因时间选择错误，想先减员再重新提交。机器人多次进入未形成有效回复的状态，仅展示推荐问题，未明确说明单位职工减员办理路径、可操作入口和注意事项，问题未闭环。",
     "补充“新增时间填报错误、参保职工减少、重新提交”等关联问法，建立单位职工减员与信息更正的分支回复；首轮先确认办理主体，再给出线上入口、办理前提和注意事项，避免连续空回复。"),
    ("案例二：灵活就业人员退休年龄咨询，回答未落到具体结论", "case_bad_2_SS202608042252101994324906.png",
     "市民提供出生年月、参保经历和“最早何时退休”等关键信息，机器人先后重复追问弹性提前或弹性延迟，未结合用户明确的“最早”诉求给出适用规则、年龄下限和办理前提。",
     "增加“出生年月、灵活就业、最早退休年龄”组合问法的优先识别；先回答法定退休年龄和弹性提前退休边界，再补充最低缴费年限、待遇领取地等影响因素，并提示以经办机构核定结果为准。"),
    ("案例三：失业金与8月缴费同时咨询，多意图未承接", "case_bad_3_SS202608041800051821394547.png",
     "市民同时咨询8月失业金领取、单位减员后何时可以领取，以及8月是否还需缴纳养老和医疗保险。机器人只围绕停保主体反复追问，未拆分失业待遇、养老缴费和医疗缴费三个问题。",
     "增加多意图拆解流程：先分别回应失业金申领条件及时间，再说明减员后养老、医疗缴费衔接方式，最后给出线上办理入口和咨询渠道；对医保事项明确提示由医保经办机构负责，避免只回答其中一项。"),
]

GOOD_CASES = [
    ("案例一：跨省转移接续，办理路径清晰完整", "case_good_1_SS202608191636571501410525.png",
     "机器人准确识别跨省转移接续事项，给出国家社会保险公共服务平台、掌上12333和电子社保卡三类办理渠道，并说明线上申请和进度查询路径，同时补充省内跨区域就业无需办理转移，信息完整且具有可执行性。"),
    ("案例二：个人姓名变更，按办理主体区分路径", "case_good_2_SS202608191638481295270544.png",
     "机器人围绕“社保改名字”区分单位职工和个体劳动者两类场景，分别提供浙江政务服务网、浙里办及线下窗口办理方式，并列出材料提示和社保易窗线上申报渠道，形成了从判断主体到选择办理渠道的完整指引。"),
    ("案例三：工伤认定，边界说明与后续指引到位", "case_good_3_SS202608152051441142811875.png",
     "机器人先明确工伤认定不属于社保中心工作范围，再提供工伤认定申请事项链接，并衔接工伤医疗费用、劳动能力鉴定和伤残待遇申领路径，既完成职责边界说明，又给出后续可执行步骤。"),
]


def img_font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Songti.ttc",
    ]
    path = next((p for p in candidates if Path(p).exists()), None)
    try:
        return ImageFont.truetype(path, size) if path else ImageFont.load_default()
    except OSError:
        return ImageFont.load_default()


def set_east_asia(run, name="Songti SC"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_run(run, size=14, bold=False, color=None):
    set_east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def replace_paragraph(p, text, size=14, bold=False, align=None, color=None):
    p.clear()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run(r, size, bold, color)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run(r)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, size=16, bold=True)
    p._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))
    return p


def add_case(doc, item, kind):
    title, image, body = item[:3]
    fill = "E8A33B" if kind == "bad" else "53B495"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run(r, bold=True, color=(255, 255, 255))
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)
    ip = doc.add_paragraph()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(str(RUNTIME / image), width=Inches(6.45))
    ip.paragraph_format.space_after = Pt(4)
    label = "问题表现：" if kind == "bad" else "优秀表现："
    add_body(doc, label + body, bold_prefix=label)
    if kind == "bad":
        add_body(doc, "优化动作：" + item[3], bold_prefix="优化动作：")


def prepare_images():
    # Use full-width screenshot crops similar to the Dancheng template.
    src = Image.open(RUNTIME / "base_cards_current.png")
    overview = src.crop((250, 250, 1395, 690)).resize((1100, 460))
    overview.save(RUNTIME / "dancheng_style_overview.png")

    src = Image.open(RUNTIME / "base_charts_current.png")
    reception = src.crop((250, 220, 1395, 475)).resize((1100, 245))
    reception.save(RUNTIME / "dancheng_style_reception.png")

    draw_chart(RUNTIME / "dancheng_style_consultation.png", CONSULTATION_DAILY + [CONSULTATION_LAST3_TOTAL],
               (77, 180, 181), "咨询量分布",
               [f"08-{i:02d}" for i in range(1, 25)] + ["08/25-27"],
               f"8月25日-8月27日合计 {CONSULTATION_LAST3_TOTAL:,} 次")


def draw_chart(path, values, color, title, labels, note):
    w, h = 1100, 350
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.text((32, 18), title, fill=(45, 45, 45), font=img_font(23, True))
    left, right, top, bottom = 65, 1060, 65, 280
    vmax = max(values) * 1.16
    for i in range(5):
        y = bottom - (bottom - top) * i / 4
        d.line((left, y, right, y), fill=(225, 230, 238), width=1)
        d.text((8, y - 9), f"{int(vmax * i / 4):,}", fill=(95, 95, 95), font=img_font(13))
    points = []
    for i, val in enumerate(values):
        x = left + (right - left) * i / (len(values) - 1)
        y = bottom - (bottom - top) * val / vmax
        points.append((x, y))
    d.line(points, fill=color, width=4, joint="curve")
    for i, (x, y) in enumerate(points):
        d.ellipse((x - 3, y - 3, x + 3, y + 3), fill=color)
        if i in (0, len(points) - 1) or len(points) <= 10:
            d.text((x - 18, y - 23), f"{values[i]:,}", fill=(50, 50, 50), font=img_font(12))
        if i % max(1, math.ceil(len(points) / 10)) == 0 or i == len(points) - 1:
            d.text((x - 22, bottom + 10), labels[i], fill=(90, 90, 90), font=img_font(12))
    d.text((65, 320), note, fill=(125, 125, 125), font=img_font(13))
    im.save(path)


def set_cell(cell, value, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value))
    set_run(r, size=10.5 if not header else 11, bold=header)
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), "D8EAF6" if header else "FFFFFF")
    cell.vertical_alignment = 1


def add_high_frequency_table(doc, after_paragraph):
    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    after_paragraph._p.addnext(page_break._p)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    r = heading.add_run("市民咨询高频问题")
    set_run(r, size=14, bold=True)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell(table.cell(0, 0), "知识库名称", True, WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(table.cell(0, 1), "会话量", True, WD_ALIGN_PARAGRAPH.CENTER)
    for name, count in TOP20:
        cells = table.add_row().cells
        set_cell(cells[0], name)
        set_cell(cells[1], f"{count:,}", align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break._p.addnext(heading._p)
    heading._p.addnext(table._tbl)


def replace_body_media(docx_path, image_paths):
    temp = docx_path.with_suffix(".media.tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as dst:
        replacements = {
            "word/media/image2.png": image_paths[0].read_bytes(),
            "word/media/image3.png": image_paths[1].read_bytes(),
            "word/media/image4.png": image_paths[2].read_bytes(),
        }
        for info in src.infolist():
            dst.writestr(info, replacements.get(info.filename, src.read(info.filename)))
    temp.replace(docx_path)


def update_heights(docx_path, image_paths):
    doc = Document(docx_path)
    for shape, path in zip(doc.inline_shapes, image_paths):
        with Image.open(path) as im:
            shape.height = int(shape.width * im.height / im.width)
    doc.save(docx_path)


def main():
    prepare_images()
    images = [
        RUNTIME / "dancheng_style_overview.png",
        RUNTIME / "dancheng_style_reception.png",
        RUNTIME / "dancheng_style_consultation.png",
    ]
    doc = Document(TEMPLATE)
    p = doc.paragraphs

    replace_paragraph(p[1], "“杭州社保”AI运营巡检报告", 24, color=(226, 88, 111), align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[3], "以下是2026年8月1日-8月27日“杭州社保”智能AI运行的详细巡检报告，我们将从咨询数据分析、会话明细分析、市民咨询高频问题、线上优化动作、待优化案例和优秀案例等方面进行梳理，以持续提升AI服务的精准度、稳定性和问题解决能力。")
    replace_paragraph(p[4], "一、咨询数据分析", 16)
    replace_paragraph(p[6], "咨询情况", 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[8], "接待量分布", 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[9], "总体咨询情况。2026年8月1日-8月27日，“杭州社保”智能机器人接待量为17,697次，咨询量为37,239次，平均对话总轮次和有效平均对话轮次均为2.1轮。")
    replace_paragraph(p[10], "2、当前AI准确率。系统侧AI准确率未形成有效数值，人工巡检正确数和错误数均为0，未打标量为37,239条；本期人工复核1,000条记录，其中有效回复666条，发现待优化问题334条，复核准确率为66.6%。")
    replace_paragraph(p[11], "3、当前咨询量分析。8月1日-8月27日接待量在67-1,039次之间波动，8月18日达到接待量峰值1,039次；咨询量集中在工作日，8月25日-8月27日合计3,654次。")
    replace_paragraph(p[12], "二、会话明细分析", 16)
    replace_paragraph(p[14], "咨询量分布", 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    replace_paragraph(p[15], "1、会话场景分析。市民咨询主要集中在人工客服、人社小灵光、问候语、网报新增回退申请表下载、业务办理、个人权益记录查询打印、个体劳动者参保登记与停保登记、跨省转移接续等事项。前20项高频问题详见下表。")
    replace_paragraph(p[16], "2、会话结构分析。本期人工复核1,000条记录，涉及556个会话；有效回复666条，待优化问题334条，覆盖260个会话，复核准确率为66.6%。问题主要表现为未形成有效回复、复杂问题未承接、多意图拆解不足及政策类问答未落到具体结论。")
    add_high_frequency_table(doc, p[16])
    replace_paragraph(p[18], "三、线上优化AI动作", 16)
    replace_paragraph(p[19], "1、围绕停保、减员、补缴、重复缴费、失业待遇及退休政策等高频复杂场景补充口语化问法、组合条件和办理分支，提升问题识别与答案覆盖。")
    replace_paragraph(p[20], "2、建立多意图拆解和上下文承接机制，先分别回应用户问题，再补充办理入口、材料要求、责任边界和下一步操作。")
    replace_paragraph(p[21], "3、对未形成有效回复、连续追问和政策边界不清场景设置兜底话术，并将334条待优化问题纳入回归复测和持续复核。")
    replace_paragraph(p[23], "四、待优化案例与线上优化动作", 16)
    replace_paragraph(p[24], "本期人工复核1,000条记录，发现待优化问题334条，覆盖260个会话；有效回复666条，复核准确率为66.6%。以下列出3例典型待优化案例。")
    replace_paragraph(p[25], "待优化问题集中在未形成有效回复、复杂问题未承接、多意图拆解不足及政策类问答未落到具体结论。")
    replace_paragraph(p[26], "优化方向：补齐高频复杂问法，完善主体识别、分支回复、兜底话术和人工衔接，并持续开展回归复测。")

    doc.add_page_break()
    add_heading(doc, "待优化案例截图")
    for i, item in enumerate(BAD_CASES):
        if i:
            doc.add_page_break()
        add_case(doc, item, "bad")

    doc.add_page_break()
    add_heading(doc, "五、本期优秀案例")
    add_body(doc, "本期优秀案例体现出较好的场景识别、边界说明、办理路径指引和问题闭环能力，可沉淀为“识别场景—说明结论—提供入口—提示下一步”的通用答复结构。")
    for i, item in enumerate(GOOD_CASES):
        if i:
            doc.add_page_break()
        add_case(doc, item, "good")

    add_heading(doc, "六、后续优化重点")
    add_body(doc, "1. 聚焦未形成有效回复的问题，优先补齐停保减员、重复缴费、失业待遇、医保衔接和退休政策等高频复杂场景。")
    add_body(doc, "2. 完善多轮对话状态记忆，避免重复追问；对同一用户的多个问题进行拆分回答，确保每个诉求都有明确结论。")
    add_body(doc, "3. 将跨省转移、个人信息变更、工伤认定等优秀答复结构沉淀为可复用话术模板，提升相近问题的服务一致性。")

    doc.save(OUTPUT)
    replace_body_media(OUTPUT, images)
    update_heights(OUTPUT, images)
    replace_body_media(OUTPUT, images)
    print(OUTPUT)


if __name__ == "__main__":
    main()
